"""
ExcelFusion - versión web (Streamlit)
=====================================
Subí varios Excel/CSV, mirá el dashboard con KPIs y descargá un reporte listo.

Para correrla en tu PC:
    pip install -r requirements.txt
    streamlit run app.py
"""
import io
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st

import marca
from consolidar_excels import EXTENSIONES, armar_resumen, formatear_hoja, leer_archivo, limpiar_nombre
from dashboard import PERIODOS, calcular_kpis, detectar_columnas, elegir_defaults, filtrar, ranking, serie_tiempo
from excel_dashboard import agregar_dashboard
from limpieza import limpiar
from mapeo import (COL_EN, COL_NUEVO, COL_ORIGINAL, aplicar_mapeo, hay_columnas_desparejas,
                   sugerir_mapeo, tabla_columnas)

COLUMNAS_LOG = ["archivo", "hoja", "estado", "detalle"]


def leer_subidos(archivos, quitar_totales=True):
    """Lee los archivos subidos. Devuelve (lista de tablas, registro)."""
    registro, frames = [], []
    with tempfile.TemporaryDirectory() as tmp:
        for i, archivo in enumerate(archivos):
            ruta = Path(tmp) / Path(archivo.name).name
            if ruta.exists():  # dos archivos con el mismo nombre
                ruta = Path(tmp) / f"{ruta.stem}_{i}{ruta.suffix}"
            ruta.write_bytes(archivo.getvalue())
            frames += leer_archivo(ruta, registro, quitar_totales)
    return frames, registro


def unir(frames, registro):
    """Junta las tablas en una sola y avisa de las columnas que no están en todos los archivos."""
    registro = list(registro)
    consolidado = pd.concat(frames, ignore_index=True, sort=False)
    for col in consolidado.columns:
        if col in ("archivo_origen", "hoja_origen"):
            continue
        faltan = sorted({f["archivo_origen"].iat[0] for f in frames if col not in f.columns})
        if faltan:
            registro.append(("-", "-", "AVISO", f"Columna '{col}' falta en: {', '.join(faltan)}"))
    return consolidado, pd.DataFrame(registro, columns=COLUMNAS_LOG)


def cargar(archivos, quitar_totales=True, mapeo=None):
    """Atajo: lee, aplica el mapeo (el sugerido si no se da uno) y une. Devuelve (consolidado | None, registro)."""
    frames, registro = leer_subidos(archivos, quitar_totales)
    if not frames:
        return None, pd.DataFrame(registro, columns=COLUMNAS_LOG)
    return unir(aplicar_mapeo(frames, mapeo or sugerir_mapeo(frames)), registro)


def exportar_excel(consolidado, resumen, log, dashboard=None, limpieza=None) -> bytes:
    """Arma el reporte en memoria y devuelve los bytes del .xlsx.
    Hojas: Dashboard (con gráficos de Excel, si hay datos numéricos), Consolidado, Resumen, Limpieza (si se
    limpió algo) y Registro. 'limpieza' es (resumen de la limpieza, filas eliminadas)."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl", datetime_format="DD/MM/YYYY") as xl:
        consolidado.to_excel(xl, sheet_name="Consolidado", index=False)
        resumen.to_excel(xl, sheet_name="Resumen", index=False)
        if limpieza is not None:
            resumen_limpieza, eliminadas = limpieza
            resumen_limpieza.to_excel(xl, sheet_name="Limpieza", index=False)
            if len(eliminadas):
                eliminadas.to_excel(xl, sheet_name="Limpieza", index=False, startrow=len(resumen_limpieza) + 3)
        log.to_excel(xl, sheet_name="Registro", index=False)
        for hoja in xl.book.worksheets:
            formatear_hoja(hoja)
            if hoja.title == "Limpieza":
                hoja.auto_filter.ref = None  # son dos tablas: sin filtro global
                if len(limpieza[1]):
                    from openpyxl.styles import Font, PatternFill
                    fila = len(limpieza[0]) + 3
                    hoja.cell(fila, 1, "Filas eliminadas (revisalas):").font = Font(bold=True)
                    for celda in hoja[fila + 1]:
                        if celda.value is not None:
                            celda.font = Font(bold=True)
                            celda.fill = PatternFill("solid", fgColor="D9E1F2")
        if dashboard:
            agregar_dashboard(xl.book, dashboard)
    return buffer.getvalue()


def numero(valor) -> str:
    """12345.678 -> '12,345.68'."""
    return f"{valor:,.2f}"


PERIODO_EN = {"D": "day", "W": "week", "M": "month"}


def dominio(valores):
    """Rango del eje Y: empieza en 0 (o en el mínimo si hay negativos) y deja un 8% de aire arriba."""
    v = pd.Series(valores).dropna()
    if v.empty:
        return [0, 1]
    bajo, alto = min(0.0, float(v.min())), max(0.0, float(v.max()))
    if alto == bajo:
        alto = bajo + 1
    return [bajo * 1.08 if bajo < 0 else 0, alto * 1.08]


def grafico_ranking(rank, categoria, metrica):
    """Barras verdes ordenadas de mayor a menor. Devuelve (datos, especificación Vega-Lite)."""
    datos = pd.DataFrame({"categoria": rank[categoria].astype(str), "valor": rank[metrica]})
    spec = {
        "height": 320,
        "mark": {"type": "bar", "color": marca.VERDE, "cornerRadiusTopLeft": 2, "cornerRadiusTopRight": 2},
        "encoding": {
            "x": {"field": "categoria", "type": "nominal", "sort": "-y", "title": None,
                  "axis": {"labelAngle": 0 if len(datos) <= 6 else -40, "labelLimit": 140}},
            "y": {"field": "valor", "type": "quantitative", "title": metrica,
                  "scale": {"domain": dominio(datos["valor"])}},
            "tooltip": [{"field": "categoria", "type": "nominal", "title": categoria},
                        {"field": "valor", "type": "quantitative", "title": metrica, "format": ",.2f"}],
        },
    }
    return datos, spec


def grafico_tiempo(serie, metrica, freq):
    """Línea azul tinta con puntos, con el eje Y desde cero. Devuelve (datos, especificación)."""
    datos = pd.DataFrame({"periodo": serie.index, "valor": serie.values})
    formato = "%b %Y" if freq == "M" else "%d %b"
    spec = {
        "height": 320,
        "mark": {"type": "line", "color": marca.TINTA, "strokeWidth": 2.5,
                 "point": {"color": marca.TINTA, "size": 45}},
        "encoding": {
            "x": {"field": "periodo", "type": "temporal", "title": None, "axis": {"format": formato}},
            "y": {"field": "valor", "type": "quantitative", "title": metrica,
                  "scale": {"domain": dominio(datos["valor"])}},
            "tooltip": [{"field": "periodo", "type": "temporal", "title": PERIODO_EN[freq], "format": formato},
                        {"field": "valor", "type": "quantitative", "title": metrica, "format": ",.2f"}],
        },
    }
    return datos, spec


def mostrar_vega(datos, spec):
    try:
        st.vega_lite_chart(datos, spec, use_container_width=True)
    except TypeError:  # versiones nuevas de Streamlit sin ese parámetro
        st.vega_lite_chart(datos, spec)


def mostrar_dashboard(df, defaults):
    """Pestaña Dashboard: elegís métrica, categoría y fechas, y ves KPIs y gráficos."""
    cols = detectar_columnas(df)
    if not cols["numericas"]:
        marca.aviso(st, "No numeric columns found to build KPIs. / No hay columnas numéricas para calcular KPIs.")
        return

    c1, c2, c3, c4 = st.columns(4)
    metrica = c1.selectbox("Metric / Métrica", cols["numericas"])
    categoria = c2.selectbox("Category / Categoría", cols["categoricas"]) if cols["categoricas"] else None
    fecha = c3.selectbox("Date / Fecha", cols["fechas"]) if cols["fechas"] else None
    freq = defaults["freq"]
    if fecha:
        opciones = list(PERIODOS)
        freq = c4.selectbox("Period / Período", opciones, index=opciones.index(freq),
                            format_func=lambda x: f"{PERIODOS[x]} / {PERIODO_EN[x].capitalize()}")

    valores, desde, hasta, todos = None, None, None, []
    with st.expander("Filters / Filtros"):
        top = st.slider("Top N", min_value=3, max_value=20, value=10)
        if categoria:
            todos = sorted(df[categoria].dropna().astype(str).unique())
            valores = st.multiselect(f"{categoria}", todos, default=todos)
        if fecha and df[fecha].notna().any():
            minimo, maximo = df[fecha].min().date(), df[fecha].max().date()
            rango = st.date_input("Date range / Rango de fechas", value=(minimo, maximo),
                                  min_value=minimo, max_value=maximo)
            if isinstance(rango, (tuple, list)) and len(rango) == 2 and tuple(rango) != (minimo, maximo):
                desde, hasta = rango

    datos = filtrar(df, categoria, valores, fecha, desde, hasta)
    if datos.empty:
        st.warning("No data for the selected filters. / No hay datos con los filtros elegidos.")
        return

    partes = []
    if categoria and valores is not None and len(valores) < len(todos):
        partes.append(f"{categoria}: " + ", ".join(valores[:4]) + ("…" if len(valores) > 4 else ""))
    if desde is not None:
        partes.append(f"fechas {desde:%d/%m/%Y} a {hasta:%d/%m/%Y}")

    k = calcular_kpis(datos, metrica, categoria, fecha, freq)
    m = st.columns(5)
    m[0].metric(f"Total {metrica}", numero(k["total"]))
    m[1].metric("Average / Promedio", numero(k["promedio"]))
    m[2].metric("Records / Registros", f"{k['registros']:,}")
    if k["lider"] is not None:
        m[3].metric(f"Top {categoria} ({k['lider_pct']:.1f}%)", k["lider"])
    if k["mejor_periodo"] is not None:
        # En semanas se muestra solo la fecha de inicio: "31/08/2026" en vez de "Semana del 31/08/2026"
        valor = k["mejor_periodo"].replace("Semana del ", "")
        rotulo = "Best week (from)" if freq == "W" else f"Best {PERIODO_EN[freq]}"
        m[4].metric(rotulo, valor, help=f"Mejor período / Best period. Total: {numero(k['mejor_periodo_valor'])}")

    rank, serie = None, None
    g1, g2 = st.columns(2)
    if categoria:
        rank = ranking(datos, categoria, metrica, top)
        with g1:
            st.subheader(f"{metrica} by {categoria}")
            mostrar_vega(*grafico_ranking(rank, categoria, metrica))
            st.dataframe(rank)
    if fecha:
        serie = serie_tiempo(datos, fecha, metrica, freq)
        with g2:
            if len(serie):
                st.subheader(f"{metrica} by {PERIODO_EN[freq]}")
                mostrar_vega(*grafico_tiempo(serie, metrica, freq))
                if freq != "D":
                    st.caption("The first and last periods may be incomplete. / "
                               "El primer y el último período pueden estar incompletos.")
            else:
                marca.aviso(st, "No valid dates to build the time chart. / No hay fechas válidas para el gráfico.")

    # Lo que se ve en pantalla es lo que se descarga: el Excel arma su hoja Dashboard con esta configuración
    return {"metrica": metrica, "categoria": categoria, "fecha": fecha, "freq": freq, "kpis": k,
            "ranking": rank if rank is not None else pd.DataFrame(), "serie": serie, "filtros": "; ".join(partes)}


def contenido():
    """Cargador, mapeo de columnas, KPIs, dashboard y descarga (todo lo que va debajo de la cabecera)."""
    subidos = st.file_uploader(
        "Upload your files / Subí tus archivos",
        type=[e.lstrip(".") for e in sorted(EXTENSIONES)],
        accept_multiple_files=True,
    )
    if not subidos:
        marca.aviso(st, "Upload two or more files to start. / Subí dos o más archivos para empezar.")
        return

    with st.expander("Options / Opciones"):
        quitar_totales = st.checkbox(
            "Skip total rows (Total, Subtotal...) / Omitir filas de totales (Total, Subtotal...)", value=True)

    # Leer los archivos es lo más lento: se hace una sola vez por cada conjunto de archivos y opciones
    firma = (tuple((a.name, a.size) for a in subidos), quitar_totales)
    if st.session_state.get("firma") != firma:
        st.session_state["firma"] = firma
        st.session_state["frames"], st.session_state["registro"] = leer_subidos(subidos, quitar_totales)
    frames, registro = st.session_state["frames"], st.session_state["registro"]
    if not frames:
        st.error("No file could be read. / No se pudo leer ningún archivo.")
        st.dataframe(pd.DataFrame(registro, columns=COLUMNAS_LOG))
        return

    # Mapeo de columnas: sugerido automáticamente, editable por la persona
    with st.expander("Column mapping / Mapeo de columnas", expanded=hay_columnas_desparejas(frames)):
        st.caption("Columns named differently in each file? Rename them to the same name to merge them. "
                   "Suggestions are pre-filled: review them. / ¿Columnas con nombres distintos en cada archivo? "
                   "Poneles el mismo nombre para unirlas. Las sugerencias ya vienen cargadas: revisalas.")
        editada = st.data_editor(tabla_columnas(frames), hide_index=True, disabled=[COL_ORIGINAL, COL_EN],
                                 key=f"mapeo_{abs(hash(firma))}")
    mapeo = {}
    for original, nuevo in zip(editada[COL_ORIGINAL], editada[COL_NUEVO]):
        mapeo[original] = (limpiar_nombre(nuevo) if isinstance(nuevo, str) else "") or original
    registro_final = list(registro) + [("-", "-", "INFO", f"Columna '{o}' unificada como '{n}'")
                                       for o, n in mapeo.items() if o != n]
    consolidado, log = unir(aplicar_mapeo(frames, mapeo), registro_final)

    # Limpieza de datos: cada cambio queda registrado y las filas eliminadas no se pierden
    with st.expander("Data cleaning / Limpieza de datos"):
        st.caption("Nothing disappears silently: every change is listed in the Registro and removed rows get their own "
                   "sheet. / Nada desaparece en silencio: cada cambio queda en el Registro y las filas eliminadas "
                   "van a su propia hoja.")
        recortar = st.checkbox("Trim extra spaces / Quitar espacios sobrantes", value=True)
        variantes = st.checkbox("Unify spelling variants (Café, CAFÉ, cafe become one) / "
                                "Unificar variantes de escritura (Café, CAFÉ y cafe pasan a ser una)", value=True)
        duplicados = st.selectbox(
            "Remove duplicate rows / Eliminar filas duplicadas", ["no", "archivo", "todos"],
            format_func=lambda x: {"no": "No", "archivo": "Only inside each file / Solo dentro de cada archivo",
                                   "todos": "Across all files / Entre todos los archivos"}[x])
        faltantes = st.checkbox("Flag rows with missing data (adds a column) / "
                                "Marcar filas con datos faltantes (agrega una columna)", value=False)
    consolidado, eliminadas, resumen_limpieza = limpiar(consolidado, recortar, variantes, duplicados, faltantes)
    cambios = resumen_limpieza[resumen_limpieza["Cantidad"] > 0]
    if len(cambios):
        marca.aviso(st, "Cleaning applied / Limpieza aplicada: " +
                    " · ".join(f"{a} {n}" for a, n, _ in cambios.itertuples(index=False)))
        log = pd.concat([log, pd.DataFrame(
            [("-", "-", "INFO", f"Limpieza: {a} = {n}. {d}") for a, n, d in cambios.itertuples(index=False)],
            columns=COLUMNAS_LOG)], ignore_index=True)

    # Columnas de texto disponibles para agrupar el resumen
    opciones = ["archivo_origen"] + [
        c for c in consolidado.select_dtypes(exclude=["number", "datetime"]).columns
        if c not in ("archivo_origen", "hoja_origen")
    ]
    agrupar = st.selectbox("Group summary by / Agrupar resumen por", opciones)
    resumen = armar_resumen(consolidado, agrupar)
    defaults = elegir_defaults(consolidado)

    c1, c2, c3 = st.columns(3)
    c1.metric("Files / Archivos", len(subidos))
    c2.metric("Rows / Filas", len(consolidado))
    c3.metric("Columns / Columnas", len(consolidado.columns))

    zona_descarga = st.container()  # el botón queda acá arriba, pero se completa al final (ya con el dashboard elegido)

    tab0, tab1, tab2, tab3 = st.tabs(["Dashboard", "Resumen", "Consolidado", "Registro"])
    with tab0:
        config = mostrar_dashboard(consolidado, defaults)
    with tab1:
        st.dataframe(resumen)
    with tab2:
        st.dataframe(consolidado)
    with tab3:
        st.dataframe(log)

    zona_descarga.download_button(
        "⬇️ Download report (.xlsx) / Descargar reporte",
        data=exportar_excel(consolidado, resumen, log, config, (resumen_limpieza, eliminadas) if len(cambios) else None),
        file_name="reporte_consolidado.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def main():
    st.set_page_config(page_title="ExcelFusion: many files, one report", page_icon=marca.icono_pagina(), layout="wide")
    marca.aplicar_estilo(st)
    marca.cabecera(st)
    contenido()
    marca.pie(st)


if __name__ == "__main__":
    main()
