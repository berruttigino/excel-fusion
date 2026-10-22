"""
ExcelFusion - hoja "Dashboard" del Excel descargable
====================================================
Arma, dentro del libro de Excel, una hoja con los indicadores clave (como celdas combinadas con
borde verde, igual que la marca) y dos gráficos NATIVOS de Excel: ranking por categoría y
evolución en el tiempo. Los gráficos se pueden editar en Excel porque leen los datos que quedan
en la propia hoja, debajo de los gráficos.
"""
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties

TINTA, VERDE, MENTA, GRIS, LINEA = "14213D", "217346", "E3F1E8", "4A5668", "D8DED8"
PERIODO_ES = {"D": "día", "W": "semana", "M": "mes"}
FORMATO_FECHA = {"D": "%d/%m/%y", "W": "%d/%m/%y", "M": "%m/%Y"}
FILA_DATOS = 27          # encabezados de las tablas de datos; los gráficos quedan arriba
ANCHO_GRAFICO, ALTO_GRAFICO = 13.4, 8.6


def _tarjetas(k, metrica, categoria, freq):
    """Indicadores a mostrar: (rótulo, valor, formato de número, tamaño de letra)."""
    tarjetas = [
        (f"Total {metrica}", k["total"], "#,##0.00", 20),
        ("Promedio por registro", k["promedio"], "#,##0.00", 20),
        ("Registros", k["registros"], "#,##0", 20),
    ]
    if k["lider"] is not None:
        tarjetas.append((f"Líder en {categoria} ({k['lider_pct']:.1f}%)", k["lider"], "General", 15))
    if k["mejor_periodo"] is not None:
        valor = k["mejor_periodo"].replace("Semana del ", "")
        etiqueta = "Mejor semana (desde)" if freq == "W" else f"Mejor {PERIODO_ES[freq]}"
        tarjetas.append((etiqueta, valor, "General", 15))
    return tarjetas


def _estilo_grafico(grafico, titulo, titulo_y, valores):
    grafico.title = titulo
    grafico.legend = None
    grafico.width, grafico.height = ANCHO_GRAFICO, ALTO_GRAFICO
    grafico.y_axis.title = titulo_y
    grafico.y_axis.number_format = "#,##0"
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E3E6E3", w=6350))
    if min(valores) >= 0:
        grafico.y_axis.scaling.min = 0


def _encabezado(celda):
    celda.font = Font(bold=True, color=TINTA)
    celda.fill = PatternFill("solid", fgColor=MENTA)
    celda.border = Border(bottom=Side(style="thin", color=VERDE))


def agregar_dashboard(wb, cfg: dict, posicion: int = 0):
    """Crea la hoja 'Dashboard'. cfg: metrica, categoria, fecha, freq, kpis, ranking, serie, filtros."""
    metrica, categoria, fecha, freq = cfg["metrica"], cfg["categoria"], cfg["fecha"], cfg["freq"]
    ranking, serie = cfg["ranking"], cfg["serie"]

    ws = wb.create_sheet("Dashboard", posicion)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = VERDE
    # Para imprimir o guardar como PDF: horizontal y ajustado al ancho de una página
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 14

    # --- encabezado ---
    ws["A1"] = f"Dashboard: {metrica}"
    ws["A1"].font = Font(size=18, bold=True, color=TINTA)
    ws.row_dimensions[1].height = 30
    partes = [f"Categoría: {categoria}" if categoria else None,
              f"Fecha: {fecha} (por {PERIODO_ES[freq]})" if fecha else None,
              f"Filtros: {cfg['filtros'] or 'ninguno'}"]
    ws["A2"] = "  ·  ".join(p for p in partes if p)
    ws["A2"].font = Font(color=GRIS)
    ws["A3"] = "Generado con ExcelFusion  ·  excelfusion.streamlit.app"
    ws["A3"].font = Font(size=9, color=GRIS)

    # --- indicadores: celdas combinadas con el borde verde de selección ---
    borde = Side(style="thick", color=VERDE)
    for i, (rotulo, valor, formato, tamano) in enumerate(_tarjetas(cfg["kpis"], metrica, categoria, freq)):
        c1, c2 = 1 + 2 * i, 2 + 2 * i
        ws.merge_cells(start_row=5, start_column=c1, end_row=5, end_column=c2)
        ws.merge_cells(start_row=6, start_column=c1, end_row=6, end_column=c2)
        etiqueta, numero = ws.cell(5, c1, rotulo), ws.cell(6, c1, valor)
        etiqueta.font = Font(size=10, color=GRIS)
        etiqueta.alignment = Alignment(horizontal="left", vertical="bottom", wrap_text=True, indent=1)
        numero.font = Font(size=tamano, bold=True, color=TINTA)
        numero.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        numero.number_format = formato
        etiqueta.border = Border(left=borde)
        numero.border = Border(left=borde)
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 34

    # --- tablas de datos (de aquí leen los gráficos) ---
    ws.cell(FILA_DATOS - 1, 1, "Datos de los gráficos").font = Font(bold=True, color=GRIS)
    fila0 = FILA_DATOS + 1
    graficos = []

    if categoria and len(ranking):
        for j, titulo in enumerate([categoria, metrica, "% participación"], 1):
            _encabezado(ws.cell(FILA_DATOS, j, titulo))
        for i, fila in enumerate(ranking.itertuples(index=False)):
            ws.cell(fila0 + i, 1, str(fila[0]))
            ws.cell(fila0 + i, 2, float(fila[1])).number_format = "#,##0.00"
            ws.cell(fila0 + i, 3, float(fila[2]) / 100).number_format = "0.0%"
        n = len(ranking)
        barras = BarChart()
        barras.type, barras.gapWidth = "col", 60
        barras.add_data(Reference(ws, min_col=2, min_row=FILA_DATOS, max_row=fila0 + n - 1), titles_from_data=True)
        barras.set_categories(Reference(ws, min_col=1, min_row=fila0, max_row=fila0 + n - 1))
        _estilo_grafico(barras, f"{metrica} por {categoria}", metrica, ranking.iloc[:, 1])
        serie_b = barras.series[0]
        serie_b.graphicalProperties.solidFill = VERDE
        serie_b.graphicalProperties.line.solidFill = VERDE
        barras.dataLabels = DataLabelList()
        barras.dataLabels.showVal = True
        for atributo in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(barras.dataLabels, atributo, False)
        graficos.append(barras)

    if fecha and serie is not None and len(serie):
        for j, titulo in enumerate([f"{PERIODO_ES[freq].capitalize()}", metrica], 6):
            _encabezado(ws.cell(FILA_DATOS, j, titulo))
        for i, (momento, valor) in enumerate(serie.items()):
            ws.cell(fila0 + i, 6, momento.strftime(FORMATO_FECHA[freq]))
            ws.cell(fila0 + i, 7, float(valor)).number_format = "#,##0.00"
        n = len(serie)
        lineas = LineChart()
        lineas.add_data(Reference(ws, min_col=7, min_row=FILA_DATOS, max_row=fila0 + n - 1), titles_from_data=True)
        lineas.set_categories(Reference(ws, min_col=6, min_row=fila0, max_row=fila0 + n - 1))
        _estilo_grafico(lineas, f"{metrica} por {PERIODO_ES[freq]}", metrica, serie)
        s = lineas.series[0]
        s.graphicalProperties.line.solidFill = TINTA
        s.graphicalProperties.line.width = 28575
        s.marker = Marker(symbol="circle", size=7)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=TINTA, ln=LineProperties(solidFill=TINTA))
        s.smooth = False
        graficos.append(lineas)

    # --- ubicación: dos gráficos lado a lado debajo de los indicadores ---
    for grafico, ancla in zip(graficos, ("A8", "F8")):
        ws.add_chart(grafico, ancla)
    for hoja in wb.worksheets:  # una sola hoja seleccionada, o Excel abriría las hojas "agrupadas"
        hoja.sheet_view.tabSelected = hoja is ws
    wb.active = wb.worksheets.index(ws)
